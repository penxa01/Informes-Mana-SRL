"""Sección OTROS — Thalamus, Ventas Mensuales, Informe IIBB, Actualizar Precios.

Los 4 flujos son los mismos de `main_legacy` (ver guideline/MainWindow.md), con los
campos ahora inline en la página en vez de diálogos modales sueltos.
"""
import os
from datetime import date, datetime

import numpy as np
import openpyxl
import pandas as pd
from dbfread import DBF
from openpyxl.utils.dataframe import dataframe_to_rows

import PyQt5.QtWidgets as widget

import src.controllers.others as others
import src.controllers.ProgressBar as ProgressBar
from src.controllers import config
from src.controllers import dbf_io
from src.controllers.others import buscaArtic
from src.ui.paginas.base import PaginaBase

_MESES = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio",
          "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"]

_FAMILIAS = ['Helix', 'Rimula', 'Gadus', 'Omala', 'Tellus', 'Advance', 'Corena',
             'Coolant', 'Diala', 'Raizen', 'Spirax', 'Refrigeration', 'Tonna', 'Turbo']

_COD_VENDEDOR = {"1": "ROMULO SANCHEZ", "3": "MIGUEL ORTUnO", "4": "JOSE BANCO",
                 "5": "DIEGO CALDERON", "6": "MATIAS BARATA", "9": "MARTIN MOYA"}

_CODIGOS_PROV = {
    "LA RIOJA": "3580", "BUENOS AIRES": "3579", "CAPITAL FEDERAL": "3578",
    "CATAMARCA": "3577", "CHACO": "3576", "CHUBUT": "3575", "CORDOBA": "3574",
    "CORRIENTES": "3573", "ENTRE RIOS": "3572", "FORMOSA": "3571", "JUJUY": "3570",
    "LA PAMPA": "3569", "MENDOZA": "3568", "MISIONES": "3567", "NEUQUEN": "3566",
    "RIO NEGRO": "3565", "SALTA": "3564", "SAN JUAN": "3563", "SAN LUIS": "3562",
    "SANTA CRUZ": "3561", "SANTA FE": "3560", "SANTIAGO DEL ESTERO": "3559",
    "TIERRA DEL FUEGO": "3558", "TUCUMAN": "3557",
}

_SUBSEGMENTOS = [
    "AGRO", "CEMENTERA/CANTERA", "CONSTRUCCIÓN", "INDUSTRIAS GENERALES",
    "METALES/CAUCHERAS", "MINERÍA", "TRANSPORTE", "OIL & GAS", "REVENTA", "ENTIDADES",
    "ENERGIA", "CONSUMO MASIVO", "LUBRICENTRO", "CASA DE REPUESTO", "CONCESIONARIOS",
    "ESTACIONES DE SERVICIOS", "TALLER", "SERVICIOS", "BARCAZA", "CARGUEROS",
    "CONTAINEROS", "CRUCEROS", "PESQUERO", "REMOLCADOR", "NAVAL", "TANQUEROS",
    "ARENERO", "DRAGAS", "FRIGORIFICOS", "CASA DE REPUESTOS MOTOS",
]


class PaginaOtros(PaginaBase):
    necesita_ventas = True

    def __init__(self, app):
        super().__init__(app)
        self._carpeta_mensual = None

        cont = widget.QWidget()
        v = widget.QVBoxLayout(cont)
        v.setContentsMargins(16, 16, 16, 16)
        v.setSpacing(14)
        v.addWidget(self._card_thalamus())
        v.addWidget(self._card_mensual())
        v.addWidget(self._card_iibb())
        v.addWidget(self._card_precios())
        v.addStretch(1)

        scroll = widget.QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setWidget(cont)
        raiz = widget.QVBoxLayout(self)
        raiz.setContentsMargins(0, 0, 0, 0)
        raiz.addWidget(scroll)

    def on_ventas_listas(self, df):
        for b in (self._btn_thalamus, self._btn_iibb):
            b.setEnabled(True)

    # ── tarjetas ─────────────────────────────────────────────────────────────
    def _card(self, titulo):
        gb = widget.QGroupBox(titulo)
        lay = widget.QHBoxLayout(gb)
        lay.setSpacing(8)
        return gb, lay

    def _card_thalamus(self):
        gb, lay = self._card("Thalamus — reporte diario de ventas (Shell)")
        lay.addWidget(widget.QLabel("Fecha:"))
        self._in_fecha = widget.QLineEdit()
        self._in_fecha.setPlaceholderText("dd, dd/mm o dd/mm/aaaa")
        self._in_fecha.setMaximumWidth(160)
        lay.addWidget(self._in_fecha)
        lay.addStretch(1)
        self._btn_thalamus = widget.QPushButton("Generar")
        self._btn_thalamus.setEnabled(False)
        self._btn_thalamus.clicked.connect(self._generar_thalamus)
        lay.addWidget(self._btn_thalamus)
        return gb

    def _card_mensual(self):
        gb, lay = self._card("Ventas Mensuales — consolidar los Thalamus diarios del mes")
        self._lbl_carpeta = widget.QLabel("(sin carpeta seleccionada)")
        self._lbl_carpeta.setStyleSheet("color: #6c757d;")
        btn_carpeta = widget.QPushButton("Elegir carpeta…")
        btn_carpeta.clicked.connect(self._elegir_carpeta)
        lay.addWidget(btn_carpeta)
        lay.addWidget(self._lbl_carpeta, 1)
        self._btn_mensual = widget.QPushButton("Consolidar")
        self._btn_mensual.clicked.connect(self._consolidar_mensual)
        lay.addWidget(self._btn_mensual)
        return gb

    def _card_iibb(self):
        gb, lay = self._card("Informe IIBB — retenciones de Ingresos Brutos (AFIP)")
        lay.addWidget(widget.QLabel("Mes:"))
        self._cb_mes = widget.QComboBox()
        self._cb_mes.addItems(_MESES)
        self._cb_mes.setCurrentIndex(date.today().month - 1)
        lay.addWidget(self._cb_mes)
        lay.addWidget(widget.QLabel("Año:"))
        self._in_anio = widget.QLineEdit(str(date.today().year))
        self._in_anio.setMaximumWidth(70)
        lay.addWidget(self._in_anio)
        lay.addStretch(1)
        self._btn_iibb = widget.QPushButton("Generar")
        self._btn_iibb.setEnabled(False)
        self._btn_iibb.clicked.connect(self._generar_iibb)
        lay.addWidget(self._btn_iibb)
        return gb

    def _card_precios(self):
        gb, lay = self._card("Actualizar Precios — abmartic.DBF de MANA + HOREB")
        self._lbl_precios = widget.QLabel("Listo.")
        self._lbl_precios.setStyleSheet("color: #6c757d;")
        lay.addWidget(self._lbl_precios, 1)
        self._btn_precios = widget.QPushButton("Actualizar MANA + HOREB")
        self._btn_precios.clicked.connect(self._actualizar_precios)
        lay.addWidget(self._btn_precios)
        return gb

    # ── 1. Thalamus ─────────────────────────────────────────────────────────
    def _parse_fecha_dia(self, txt):
        txt = txt.strip()
        hoy = date.today()
        if "/" not in txt:
            txt = f"{txt}/{hoy.month}/{hoy.year}"
        elif txt.count("/") == 1:
            txt = f"{txt}/{hoy.year}"
        return datetime.strptime(txt, "%d/%m/%Y").date()

    def _generar_thalamus(self):
        try:
            fecha_ingresada = self._parse_fecha_dia(self._in_fecha.text())
        except ValueError:
            widget.QMessageBox.warning(self, "Fecha inválida",
                                       "Ingresá una fecha en formato dd, dd/mm o dd/mm/aaaa.")
            return
        try:
            df = self.app.df_ventas
            df_filtrado = df[
                (pd.to_datetime(df["FECHA"]).dt.date == fecha_ingresada)
                & (df["MODELO"].str.match(r'^L\d'))
                & (df["CUENTA"].notna())
                & (df["TC"].isin(["FC", "NC", "FE", "CR"]))
            ]
            SKUArticDS = buscaArtic(df_filtrado)
            codigos_correctos = self.getCodigosCorrectos(df_filtrado['MODELO'], SKUArticDS)

            nuevos_datos = pd.DataFrame({
                'Systemid': df_filtrado['CUENTA'].astype(int),
                'Sale Date': pd.to_datetime(df_filtrado['FECHA']).dt.strftime('%Y/%m/%d'),
                'Transaction ID': df_filtrado.apply(
                    lambda row: "{} {} {} {}".format(
                        row['TC'], row['LETRA'], row['NROEST'], str(row['NC']).zfill(8)), axis=1),
                'Code': codigos_correctos,
                'Quantity': df_filtrado['U_VENDIDAS'].astype(int),
                'Description 1': np.where(df_filtrado['TC'] == 'FC', 'F',
                                          np.where(df_filtrado['TC'] == 'NC', 'NC', '')),
                'Amount': round(self.addPrice(df_filtrado['MODELO'], SKUArticDS).astype(float), 2),
            })
            self.verificaCliente(nuevos_datos['Systemid'])

            ruta_base = os.path.join(config.EXCEL_THALAMUS_PATH, "Template_VENTA MANA.xlsx")
            nombre_archivo = f"Template_VENTA MANA {fecha_ingresada.strftime('%d-%m')}.xlsx"
            ruta_final = os.path.join(config.EXCEL_THALAMUS_PATH, nombre_archivo)
            if not os.path.exists(ruta_base):
                widget.QMessageBox.warning(self, "Error", f"No se encontró el archivo base:\n{ruta_base}")
                return
            try:
                workbook = openpyxl.load_workbook(ruta_base)
                hoja = workbook.active
                hoja.delete_rows(2, hoja.max_row)
                for fila in dataframe_to_rows(nuevos_datos, index=False, header=False):
                    hoja.append(fila)
                for fila in hoja.iter_rows(min_row=2, min_col=2, max_col=2):
                    for celda in fila:
                        celda.number_format = "yyyy/mm/dd;@"
                workbook.save(ruta_final)
                widget.QMessageBox.information(self, "Éxito", f"Datos añadidos a:\n{ruta_final}")
            except PermissionError:
                widget.QMessageBox.critical(self, "Error", f"El archivo está abierto:\n{ruta_base}")
        except Exception as e:  # noqa: BLE001
            widget.QMessageBox.critical(self, "Error", str(e))

    def getCodigosCorrectos(self, modelos_series, SKUArticDS):
        codigos_resultado = []
        for modelo in modelos_series:
            matches = SKUArticDS[SKUArticDS['CODART'] == modelo]
            if not matches.empty:
                artiprov = matches['ARTIPROV'].iloc[0]
                codigo_final = artiprov if (artiprov is not None and artiprov != '' and pd.notna(artiprov)) else modelo
            else:
                codigo_final = modelo
            codigos_resultado.append(codigo_final)
        return pd.Series(codigos_resultado, index=modelos_series.index)

    def addPrice(self, modelos_series, SKUArticDS):
        try:
            precios = pd.read_excel(
                os.path.join(config.EXCEL_PRECIOS_PATH, 'Lista de precios Mana.xlsm'),
                sheet_name='Lista De Precios', header=2)
            precios = precios[["SKU", 'Producto', "Precio Por envase", "Informe  s/IVA"]]
            precios['SKU'] = precios['SKU'].astype(str).str.strip()
            precio_dict = dict(zip(precios['SKU'], precios['Informe  s/IVA']))
            precios_resultado = []
            for modelo in modelos_series:
                matches = SKUArticDS[SKUArticDS['CODART'] == modelo]
                if not matches.empty:
                    artiprov = matches['ARTIPROV'].iloc[0]
                    codigo_busqueda = artiprov if (artiprov is not None and artiprov != '' and pd.notna(artiprov)) else modelo
                else:
                    codigo_busqueda = modelo
                precios_resultado.append(precio_dict.get(codigo_busqueda, 0))
            return pd.Series(precios_resultado, index=modelos_series.index)
        except FileNotFoundError:
            return pd.Series([0] * len(modelos_series), index=modelos_series.index)
        except Exception as e:  # noqa: BLE001
            print(f"Error al procesar precios: {e}")
            return pd.Series([0] * len(modelos_series), index=modelos_series.index)

    def verificaCliente(self, numero_clientes):
        documento = pd.read_excel(os.path.join(config.EXCEL_THALAMUS_PATH, "Template_CLIENTE MANA.xlsx"))
        bdClient = pd.DataFrame(iter(DBF(os.path.join(config.DBF_BASE_PATH, 'ctaclien.DBF'), encoding=config.DBF_ENCODING)))
        no_encontrados = set(numero_clientes) - set(documento["Systemid"])
        try:
            if no_encontrados:
                filas_nuevas = []
                for cid in no_encontrados:
                    cliente = bdClient[bdClient['CNUMERO'] == cid]
                    if cliente.empty:
                        raise ValueError(f"No se encontró el cliente con Systemid: {cid}")
                    cuit = cliente["CUIT"].values[0].replace("-", "")
                    filas_nuevas.append({
                        "Systemid": cid,
                        "FiscalID": int(cuit) if cuit.isdigit() else "",
                        "FiscalName": cliente["CNOMBRE"].values[0],
                        "Name": cliente["CNOMBRE"].values[0],
                        "Address.countryId": 1010,
                        "Address.street1": cliente["CPROVIN"].values[0],
                        "Address.street2": cliente["CDIREC"].values[0],
                        "Address.postalCode": int(cliente["CCP"].values[0]) if cliente["CCP"].values[0].isdigit() else '',
                        "Address.city": cliente["CLOCAL"].values[0],
                        "address.stateId": int(_CODIGOS_PROV.get(cliente["CPROVIN"].values[0])),
                        "email": "", "dia_visita": "", "frecuencia_visita": "", "semana_visita": "",
                        "subsegmento": self.selectSubsec(cid, cliente["CNOMBRE"].values[0]),
                        "zona": "",
                        "Vendedor": _COD_VENDEDOR[str(cliente["VENDEDOR"].values[0])],
                    })
                if filas_nuevas:
                    documento = pd.concat([documento, pd.DataFrame(filas_nuevas)], ignore_index=True)
                    documento.sort_values(by=["Systemid"], ascending=True, inplace=True)
                    others.format_CLIENTE(documento,
                                          os.path.join(config.EXCEL_THALAMUS_PATH, "Template_CLIENTE MANA.xlsx"),
                                          len(no_encontrados))
        except ValueError as e:
            widget.QMessageBox.warning(self, "Error", str(e))

    def selectSubsec(self, cid, nombre):
        dialog = widget.QDialog(self)
        dialog.setWindowTitle("CLIENTE NUEVO - Seleccionar Subsegmento")
        lay = widget.QVBoxLayout(dialog)
        lay.addWidget(widget.QLabel(f"ID: {cid}\nNombre: {nombre}", dialog))
        combo = widget.QComboBox(dialog)
        combo.addItems(_SUBSEGMENTOS)
        lay.addWidget(combo)
        bb = widget.QDialogButtonBox(widget.QDialogButtonBox.Ok | widget.QDialogButtonBox.Cancel, dialog)
        bb.accepted.connect(dialog.accept)
        bb.rejected.connect(dialog.reject)
        lay.addWidget(bb)
        return combo.currentText() if dialog.exec_() == widget.QDialog.Accepted else None

    # ── 2. Ventas Mensuales ─────────────────────────────────────────────────
    def _elegir_carpeta(self):
        d = widget.QFileDialog.getExistingDirectory(self, "Carpeta con los Template_VENTA MANA del mes")
        if d:
            self._carpeta_mensual = d
            self._lbl_carpeta.setText(d)

    def _consolidar_mensual(self):
        if not self._carpeta_mensual:
            widget.QMessageBox.warning(self, "Sin carpeta", "Elegí primero la carpeta del mes.")
            return
        try:
            fila_distribuidor = pd.DataFrame(
                [["DISTRIBUIDOR", None, None, None, None]],
                columns=["Fecha", "Comprobante", "Codigo de producto", "Familia", "Litros"])
            columnas = ["Dia / Periodo", "Comprobante", "Codigo de producto", "Familia", "Litros"]
            df_mensual = pd.concat([fila_distribuidor, pd.DataFrame([columnas], columns=columnas)], ignore_index=True)
            df_art = pd.read_excel(config.EXCEL_LITROS_PATH, sheet_name='GENERAL')

            for archivo in os.listdir(self._carpeta_mensual):
                if not archivo.startswith('Template'):
                    continue
                daily = pd.read_excel(os.path.join(self._carpeta_mensual, archivo))
                filas_nuevas = []
                for _, venta in daily.iterrows():
                    art_f = df_art[(df_art["articulo"] == venta["Code"]) | (df_art["Artiprov"] == venta["Code"])]
                    if art_f.empty:
                        widget.QMessageBox.information(
                            self, "Código no encontrado",
                            f"El código {venta['Code']} ({venta['Sale Date']}) no está en LITROS SHELL.")
                        continue
                    art = art_f.iloc[0]
                    filas_nuevas.append({
                        "Fecha": pd.to_datetime(venta['Sale Date']).strftime('%Y/%m/%d'),
                        "Comprobante": venta['Transaction ID'],
                        "Codigo de producto": venta['Code'],
                        "Familia": self.get_familia(art),
                        "Litros": self.get_litros(art, venta['Quantity']),
                    })
                if filas_nuevas:
                    df_mensual = pd.concat([df_mensual, pd.DataFrame(filas_nuevas)], ignore_index=True)

            ruta_salida = widget.QFileDialog.getSaveFileName(
                self, "Guardar Informe Mensual (poné el MES en el nombre)", "Anexo_6_MES.xlsx",
                "Archivos Excel (*.xlsx)")[0]
            if ruta_salida:
                others.generar_archivo_excel(df_mensual, ruta_salida)
        except Exception as e:  # noqa: BLE001
            widget.QMessageBox.critical(self, "Error", f"Se produjo un error:\n{e}")

    def get_familia(self, art):
        primera = str(art["desc"]).split()[0] if str(art["desc"]).strip() else ""
        return primera if primera in _FAMILIAS else 'Otro'

    def get_litros(self, art, cantidad):
        return art['codmedida'] * cantidad

    # ── 3. Informe IIBB ─────────────────────────────────────────────────────
    def _generar_iibb(self):
        anio_txt = self._in_anio.text().strip()
        if not anio_txt.isdigit():
            widget.QMessageBox.warning(self, "Año inválido", "Ingresá un año válido (ej. 2025).")
            return
        mes = self._cb_mes.currentIndex() + 1
        anio = int(anio_txt)
        try:
            bdClient = pd.DataFrame(iter(DBF(os.path.join(config.DBF_BASE_PATH, 'ctaclien.DBF'), encoding=config.DBF_ENCODING)))
            ivatemp = dbf_io.leer_dbf(os.path.join(config.DBF_BASE_PATH, 'ivatemp.DBF'), dbf_io.COLUMNAS_IVATEMP)
            ivatemp['FECHA'] = pd.to_datetime(ivatemp['FECHA'], errors='coerce')
            filtrado = ivatemp[
                (ivatemp['FECHA'].dt.month == mes) & (ivatemp['FECHA'].dt.year == anio)
                & ((ivatemp['NROEST'] == '00002') | (ivatemp['NROEST'] == '00003'))
            ]
            total = 0
            resultados = []
            for _, row in filtrado.iterrows():
                perce_1 = row['PERCE_1']
                cliente = bdClient[bdClient['CNUMERO'] == row['CUENTA']]
                cnombre = cliente['CNOMBRE'].values[0] if not cliente.empty else ''
                ingbrut = cliente['INGBRUT'].values[0] if not cliente.empty else ''
                ingbrut = "000-000000-0" if not ingbrut.strip() else ingbrut
                cuit = str(cliente['CUIT'].values[0]).replace('-', '') if not cliente.empty else ''
                cdirec = cliente['CDIREC'].values[0] if not cliente.empty else ''
                if perce_1 != 0:
                    tc = row['TC']
                    tc = "FA" if tc == "FC" else ("CA" if tc == "NC" else tc)
                    resultados.append([
                        cnombre, cuit, cdirec, ingbrut, row["FECHA"].strftime("%d/%m/%Y"), tc,
                        str(row['NC']).split()[0].zfill(12), row['NETO'], 3.00,
                        round(perce_1 * 0.83333327, 2), round(perce_1 * 0.1666667, 2),
                    ])
                    total += perce_1
            columnas = ['CNOMBRE', 'CUIT', 'CDIREC', 'INGBRUT', 'FECHA', 'TC', 'NC', 'NETO', 'PERCENT', 'PERCE_1_83', 'PERCE_1_16']
            df_res = pd.DataFrame(resultados, columns=columnas)
            output_path = os.path.join(config.OUTPUT_PATH, f'salib {mes}-{anio}.txt')
            os.makedirs(config.OUTPUT_PATH, exist_ok=True)
            df_res.to_csv(output_path, sep='\t', index=False, header=False)
            widget.QMessageBox.information(
                self, "Informe generado",
                f"Guardado en: {output_path}\n\nTotal Ingresos Brutos: ${total:,.2f}\n"
                f"Comprobantes: {len(df_res)}")
        except Exception as e:  # noqa: BLE001
            widget.QMessageBox.critical(self, "Error", str(e))

    # ── 4. Actualizar Precios ──────────────────────────────────────────────
    def _actualizar_precios(self):
        self._btn_precios.setEnabled(False)
        self._lbl_precios.setText("Actualizando precios, por favor espere…")
        self._hilo_precios = ProgressBar.ActualizarPreciosThread(
            os.path.join(config.DBF_BASE_PATH, 'abmartic.DBF'),
            os.path.join(config.DBF_HOREB_PATH, 'abmartic.DBF'),
            os.path.join(config.EXCEL_PRECIOS_PATH, 'Lista de precios Mana.xlsm'),
            os.path.join(config.EXCEL_PRECIOS_PATH, 'Lista Valvoline.xlsx'),
        )
        self._hilo_precios.progress_actualizado.connect(self._lbl_precios.setText)
        self._hilo_precios.finalizado.connect(self._on_precios_finished)
        self._hilo_precios.start()

    def _on_precios_finished(self, exito, mensaje):
        self._btn_precios.setEnabled(True)
        self._lbl_precios.setText("Listo.")
        if exito:
            widget.QMessageBox.information(self, "Actualizar Precios", mensaje)
        else:
            widget.QMessageBox.critical(self, "Error", mensaje)
